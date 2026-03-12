from flask import Flask, render_template, request, send_file, flash, redirect, url_for, jsonify
import os
from datetime import datetime
import calendar
import pandas as pd
import psycopg2
from openpyxl import load_workbook
from openpyxl.styles import Protection, Alignment

def get_base_dir():
    """
    Get the base directory where this script is located.
    This makes all paths dynamic and portable.
    """
    return os.path.dirname(os.path.abspath(__file__))

def get_path(*relative_paths):
    """
    Build a path relative to the script's directory.
    Usage: get_path('templates', 'index.html') -> script_dir/templates/index.html
    """
    return os.path.join(get_base_dir(), *relative_paths)

# Get dynamic paths
BASE_DIR = get_base_dir()
TEMPLATE_DIR = get_path('templates')
TEMPLATE_EXCEL_FILE = get_path('monthly_attendence_report_tempelate.xlsx')
GENERATED_REPORTS_DIR = get_path('generated_reports')

# Verify template directory exists
if not os.path.exists(TEMPLATE_DIR):
    raise FileNotFoundError(
        f"Template directory not found at: {TEMPLATE_DIR}\n"
        f"Base directory: {BASE_DIR}\n"
        f"Please ensure the 'templates' folder exists in the same directory as app.py"
    )

app = Flask(__name__, template_folder=TEMPLATE_DIR)
app.secret_key = 'your-secret-key-change-this-in-production'

# Database connection parameters
DB_CONFIG = {
    'host': 'aws-0-ap-south-1.pooler.supabase.com',
    'port': 6543,
    'dbname': 'postgres',
    'user': 'postgres.fhsvgeacwnnvqidyhnok',
    'password': 'Bait-Alshaar20'
}

def calculate_monthly_hours(from_date, to_date):
    """Calculate monthly hours based on number of days"""
    delta = to_date - from_date
    number_of_days = delta.days + 1
    monthly_hours = number_of_days * 8
    return monthly_hours, number_of_days

def generate_filename_from_date_range(from_date, to_date, filter_type=None):
    """Generate filename automatically based on date range"""
    month_name = calendar.month_name[from_date.month].lower()
    year = from_date.year
    base_name = f"monthly_attendance_report_{month_name}_{year}"
    
    if filter_type:
        base_name += f"_{filter_type.lower()}"
    
    filename = f"{base_name}.xlsx"
    return filename

def get_employees_list(from_date, to_date, filter_type=None):
    """
    Get list of employees based on filter type (construction/maintenance/all)
    Returns list of (employee_id, name) tuples
    """
    connection = psycopg2.connect(**DB_CONFIG)
    cursor = connection.cursor()
    
    try:
        if filter_type:
            # Filter directly by Employee.department (case-insensitive)
            query = '''
                SELECT DISTINCT
                    e.employee_id,
                    e.name
                FROM 
                    "Employee" e
                INNER JOIN 
                    "Attendance" a ON e.employee_id = a.employee_id
                WHERE 
                    a.date BETWEEN %s AND %s
                    AND UPPER(e.department) = UPPER(%s)
                ORDER BY 
                    e.name;
            '''
            cursor.execute(query, (from_date, to_date, filter_type))
        else:
            # Get all employees with attendance in date range
            query = '''
                SELECT DISTINCT
                    e.employee_id,
                    e.name
                FROM 
                    "Employee" e
                INNER JOIN 
                    "Attendance" a ON e.employee_id = a.employee_id
                WHERE 
                    a.date BETWEEN %s AND %s
                ORDER BY 
                    e.name;
            '''
            cursor.execute(query, (from_date, to_date))
        
        employees = cursor.fetchall()
        return employees
    finally:
        cursor.close()
        connection.close()

def generate_attendance_report(from_date, to_date, monthly_hours=None, selected_employees=None, filter_type=None):
    """
    Generate attendance report and return the file path
    selected_employees: list of employee_ids to include (None = all)
    filter_type: 'construction', 'maintenance', or None
    """
    # Use dynamic paths
    template_path = TEMPLATE_EXCEL_FILE
    
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template file not found at: {template_path}")
    
    # Calculate monthly hours if not provided
    if monthly_hours is None:
        monthly_hours, _ = calculate_monthly_hours(from_date, to_date)
    
    # Generate output filename
    output_filename = generate_filename_from_date_range(from_date, to_date, filter_type)
    # Ensure generated_reports directory exists
    os.makedirs(GENERATED_REPORTS_DIR, exist_ok=True)
    file_path = os.path.join(GENERATED_REPORTS_DIR, output_filename)
    
    # Load template
    try:
        workbook = load_workbook(template_path)
        template_sheet = workbook['Sheet1']
    except Exception as e:
        raise ValueError(f"Failed to load template file: {str(e)}")
    
    # Connect to database
    try:
        connection = psycopg2.connect(**DB_CONFIG)
        cursor = connection.cursor()
    except Exception as e:
        workbook.close()
        raise ValueError(f"Failed to connect to database: {str(e)}")
    
    try:
        # Build query with employee filter
        base_query = '''
            SELECT 
                e.employee_id,
                e.name,
                e.salary,
                a.date,
                a.notes,
                a.status,
                a.status_attendance,
                p.project_name,
                ap.working_hours,
                ap.overtime_hours
            FROM 
                "Employee" e
            LEFT JOIN 
                "Attendance" a ON e.employee_id = a.employee_id
            LEFT JOIN
                "Attendance_projects" ap ON a.id = ap.attendance_id
            LEFT JOIN
                projects p ON ap.project_id = p.project_id
            WHERE 
                a.date BETWEEN %s AND %s
        '''
        
        params = [from_date, to_date]
        
        # Add employee filter if specified
        if selected_employees:
            placeholders = ','.join(['%s'] * len(selected_employees))
            base_query += f' AND e.employee_id IN ({placeholders})'
            params.extend(selected_employees)
        
        # Add department filter directly on Employee.department
        if filter_type:
            base_query += ' AND UPPER(e.department) = UPPER(%s)'
            params.append(filter_type)
        
        base_query += ' ORDER BY e.employee_id, a.date;'
        
        cursor.execute(base_query, params)
        data = cursor.fetchall()
        
        if not data:
            raise ValueError(f"No attendance data found for the selected criteria")
        
        df = pd.DataFrame(data, columns=[
            'employee_id', 'name', 'salary', 'date', 'notes', 'status',
            'status_attendance', 'project_name', 'working_hours', 'overtime_hours'
        ])
        
        grouped_data = df.groupby('employee_id')
        
        for employee_id, employee_data in grouped_data:
            employee_sheet = workbook.copy_worksheet(template_sheet)
            
            first_row = employee_data.iloc[0]
            employee_sheet['B4'] = first_row['name']
            employee_sheet.title = f"{first_row['name']}"[:31]  # Excel sheet name limit
            employee_sheet.sheet_view.rightToLeft = template_sheet.sheet_view.rightToLeft
            employee_sheet['I2'] = first_row['salary']
            
            # Write date range and month name to the correct template cells
            employee_sheet['F2'] = from_date
            employee_sheet['F3'] = to_date
            employee_sheet['F4'] = calendar.month_name[from_date.month]
            
            # Apply cell protection
            try:
                for r in range(7, 50):
                    for c in ['B', 'C', 'D', 'E', 'F', 'H', 'J']:
                        try:
                            cell = employee_sheet[f"{c}{r}"]
                            cell.protection = Protection(locked=False)
                        except:
                            pass
            except:
                pass
            
            row_date = 7
            daily_grouped = employee_data.groupby('date')
            
            for date_value, day_data in daily_grouped:
                status = day_data.iloc[0]['status']
                status_attendance = day_data.iloc[0]['status_attendance']
                notes = day_data.iloc[0]['notes']
                
                employee_sheet[f'A{row_date}'] = date_value
                employee_sheet[f'J{row_date}'] = notes
                
                # Write attendance status
                if status == 'present' and status_attendance == 'Present':
                    employee_sheet[f'B{row_date}'] = 'P'
                elif status == 'present' and status_attendance == 'Weekend':
                    employee_sheet[f'B{row_date}'] = 'W'
                elif status == 'present' and status_attendance == 'Holiday-Work':
                    employee_sheet[f'B{row_date}'] = 'H'
                elif status == 'absent' and status_attendance == 'Absence without excuse':
                    employee_sheet[f'B{row_date}'] = 'AWO'
                elif status == 'absent' and status_attendance == 'Sick Leave':
                    employee_sheet[f'B{row_date}'] = 'SL'
                elif status == 'absent' and status_attendance == 'Absence with excuse':
                    employee_sheet[f'B{row_date}'] = 'A'
                elif status == 'vacation':
                    employee_sheet[f'B{row_date}'] = 'V'
                
                total_working_hours = day_data['working_hours'].sum()
                if total_working_hours == 0:
                    if status_attendance in ('Weekend', 'Holiday-Work'):
                        total_working_hours = 8
                
                if total_working_hours != 0:
                    employee_sheet[f'C{row_date}'] = total_working_hours
                
                total_overtime = day_data['overtime_hours'].sum()
                if total_overtime != 0:
                    if status_attendance == 'Present':
                        employee_sheet[f'D{row_date}'] = total_overtime
                    elif status_attendance == 'Weekend':
                        employee_sheet[f'E{row_date}'] = total_overtime
                    elif status_attendance == 'Holiday-Work':
                        employee_sheet[f'F{row_date}'] = total_overtime
                
                # Build project string
                if len(day_data) == 1:
                    project_text = day_data.iloc[0]['project_name']
                else:
                    project_parts = []
                    for _, proj_row in day_data.iterrows():
                        proj_name = proj_row['project_name'] if proj_row['project_name'] else 'Unknown'
                        proj_hours = int(proj_row['working_hours']) if proj_row['working_hours'] is not None and proj_row['working_hours'] != 0 else 0
                        project_parts.append(f"{proj_name} {proj_hours}hrs")
                    project_text = " + ".join(project_parts)
                
                employee_sheet[f'H{row_date}'] = project_text
                
                try:
                    employee_sheet[f'J{row_date}'].alignment = Alignment(wrap_text=True)
                    employee_sheet.row_dimensions[row_date].height = 32
                except:
                    pass
                
                row_date += 1
        
        # Ensure file path is valid and doesn't contain invalid characters
        try:
            # Normalize the path and ensure it's valid
            file_path = os.path.normpath(file_path)
            
            # Check for Windows path length limit (260 characters)
            if len(file_path) > 260:
                # Try to shorten the filename if path is too long
                dir_path = os.path.dirname(file_path)
                filename = os.path.basename(file_path)
                # Shorten filename if needed
                if len(filename) > 50:
                    name, ext = os.path.splitext(filename)
                    filename = name[:45] + ext
                file_path = os.path.join(dir_path, filename)
            
            # Ensure the directory exists
            os.makedirs(os.path.dirname(file_path), exist_ok=True)
            
            # Save the workbook
            workbook.save(file_path)
            workbook.close()
            return file_path, len(grouped_data)
        except OSError as e:
            workbook.close()
            error_msg = f"Failed to save file: {str(e)}"
            # Sanitize error message for display
            try:
                error_msg = error_msg.encode('ascii', 'ignore').decode('ascii')
            except:
                error_msg = "Failed to save file. Path may be too long or contain invalid characters."
            raise ValueError(error_msg)
        except Exception as e:
            workbook.close()
            raise
        
    finally:
        cursor.close()
        connection.close()

@app.route('/')
def index():
    # Get default date range (current month)
    today = datetime.now()
    from_date_default = datetime(today.year, today.month, 1).date()
    
    # Get last day of current month
    if today.month == 12:
        to_date_default = datetime(today.year, 12, 31).date()
    else:
        to_date_default = datetime(today.year, today.month + 1, 1).date() - pd.Timedelta(days=1)
    
    # Get employees list for the default date range
    try:
        employees = get_employees_list(from_date_default, to_date_default)
    except Exception as e:
        employees = []
        flash(f'Could not load employees list: {str(e)}', 'error')
    
    return render_template('index.html', 
                         employees=employees,
                         from_date_default=from_date_default.strftime('%Y-%m-%d'),
                         to_date_default=to_date_default.strftime('%Y-%m-%d'))

@app.route('/get_employees', methods=['POST'])
def get_employees():
    """AJAX endpoint to get employees list based on date range and filter"""
    try:
        from_date_str = request.json.get('from_date')
        to_date_str = request.json.get('to_date')
        filter_type = request.json.get('filter_type', '')
        
        from_date = datetime.strptime(from_date_str, '%Y-%m-%d').date()
        to_date = datetime.strptime(to_date_str, '%Y-%m-%d').date()
        
        employees = get_employees_list(from_date, to_date, filter_type if filter_type else None)
        
        return jsonify({'success': True, 'employees': [{'id': emp[0], 'name': emp[1]} for emp in employees]})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)})

@app.route('/generate', methods=['POST'])
def generate_report():
    try:
        # Get form data
        from_date_str = request.form.get('from_date')
        to_date_str = request.form.get('to_date')
        monthly_hours_str = request.form.get('monthly_hours', '').strip()
        selected_employees = request.form.getlist('selected_employees')  # Get list of selected employee IDs
        filter_type = request.form.get('filter_type', '').strip() or None
        
        # Parse dates
        from_date = datetime.strptime(from_date_str, '%Y-%m-%d').date()
        to_date = datetime.strptime(to_date_str, '%Y-%m-%d').date()
        
        if to_date < from_date:
            flash('Error: TO date must be after FROM date!', 'error')
            return redirect(url_for('index'))
        
        # Convert selected employees to integers if provided
        employee_ids = None
        if selected_employees:
            try:
                employee_ids = [int(emp_id) for emp_id in selected_employees if emp_id]
            except ValueError:
                flash('Invalid employee selection!', 'error')
                return redirect(url_for('index'))
        
        # Calculate or use provided monthly hours
        if monthly_hours_str:
            monthly_hours = float(monthly_hours_str)
        else:
            monthly_hours, num_days = calculate_monthly_hours(from_date, to_date)
            flash(f'Auto-calculated monthly hours: {monthly_hours} hours ({num_days} days × 8 hours/day)', 'info')
        
        # Generate report
        file_path, employee_count = generate_attendance_report(
            from_date, to_date, monthly_hours, employee_ids, filter_type
        )
        
        filename = os.path.basename(file_path)
        flash(f'Successfully generated report for {employee_count} employee(s)!', 'success')
        
        return render_template('success.html', 
                             filename=filename,
                             file_path=file_path,
                             employee_count=employee_count)
        
    except ValueError as e:
        flash(f'Error: {str(e)}', 'error')
        return redirect(url_for('index'))
    except Exception as e:
        # Safely handle error messages that might contain invalid characters
        error_type = type(e).__name__
        error_msg = str(e)
        
        # Sanitize error message to avoid encoding issues
        try:
            # Try to encode/decode to remove problematic characters
            error_msg_clean = error_msg.encode('utf-8', errors='ignore').decode('utf-8', errors='ignore')
            # Remove any remaining problematic characters
            error_msg_clean = ''.join(c for c in error_msg_clean if c.isprintable() or c.isspace())
            if not error_msg_clean.strip():
                error_msg_clean = "An unknown error occurred"
        except:
            error_msg_clean = "An error occurred while generating the report"
        
        # Flash error message (safe for web display)
        try:
            flash(f'Error generating report: {error_msg_clean}', 'error')
        except:
            flash('Error generating report. Please check your inputs and try again.', 'error')
        
        # Log error safely without traceback (which can cause OSError)
        try:
            import sys
            import logging
            logging.basicConfig(level=logging.ERROR)
            logger = logging.getLogger(__name__)
            logger.error(f"Report generation failed: {error_type}: {error_msg_clean}")
        except:
            # If logging fails, silently continue
            pass
        
        return redirect(url_for('index'))

@app.route('/download/<filename>')
def download_file(filename):
    # Use dynamic path
    file_path = os.path.join(GENERATED_REPORTS_DIR, filename)
    
    if os.path.exists(file_path):
        return send_file(file_path, as_attachment=True)
    else:
        flash('File not found!', 'error')
        return redirect(url_for('index'))

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)
